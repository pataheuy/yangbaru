import openpyxl
import os
import re

folder = r'C:\Users\ASSYIFA\OneDrive\Documents\GitHub\yangbaru\cadanganputra-main\aplikasipemira\data siswa'

def normalize(nama):
    """Normalisasi nama: uppercase, hapus tanda hubung, spasi berulang."""
    nama = nama.strip().upper()
    nama = nama.replace('-', ' ')   # AL-FARIDZI -> AL FARIDZI
    nama = re.sub(r'\s+', ' ', nama)
    return nama

# =============================================
# 1. Kumpulkan semua NIS dari Daftar NIS siswa
# =============================================
wb_nis = openpyxl.load_workbook(os.path.join(folder, 'Daftar NIS siswa.xlsx'))

# Map: nama_normalized -> NIS
nis_map = {}

for sheet_name in wb_nis.sheetnames:
    ws = wb_nis[sheet_name]
    for row in ws.iter_rows(min_row=6, values_only=True):
        no   = row[0]
        nis  = row[1]
        nama = row[2]
        # Hanya baris data valid (nomor urut angka, nis angka, nama string)
        if isinstance(no, int) and isinstance(nis, (int, float)) and isinstance(nama, str):
            key = normalize(nama)
            nis_val = str(int(nis))
            if key in nis_map and nis_map[key] != nis_val:
                print(f"[DUPLIKAT BERBEDA] {key}: {nis_map[key]} vs {nis_val}")
            nis_map[key] = nis_val

print(f"Total siswa di Daftar NIS: {len(nis_map)}")

# =============================================
# 2. Isi kolom Password di Data siswa.xlsx
# =============================================
wb_data = openpyxl.load_workbook(os.path.join(folder, 'Data siswa.xlsx'))

tidak_ditemukan = []
total_diisi = 0

for sheet_name in wb_data.sheetnames:
    ws = wb_data[sheet_name]
    for row in ws.iter_rows(min_row=2):  # skip header
        cell_nama     = row[1]  # kolom B
        cell_password = row[2]  # kolom C
        if cell_nama.value is None:
            continue
        nama_key = normalize(str(cell_nama.value))
        if nama_key in nis_map:
            cell_password.value = nis_map[nama_key]
            total_diisi += 1
        else:
            tidak_ditemukan.append(f"[{sheet_name}] {cell_nama.value}")

# =============================================
# 3. Simpan
# =============================================
output_path = os.path.join(folder, 'Data siswa.xlsx')
wb_data.save(output_path)

print(f"\nSelesai! {total_diisi} password berhasil diisi.")
if tidak_ditemukan:
    print(f"\n{len(tidak_ditemukan)} siswa TIDAK ditemukan di Daftar NIS:")
    for n in tidak_ditemukan:
        print(" -", n)
else:
    print("Semua siswa berhasil dicocokkan!")
